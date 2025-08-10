# utils/docx.py
import os
import logging
from docx import Document
from PyQt5.QtWidgets import QMessageBox
from PyQt5.QtCore import QCoreApplication

def populate_docx_template(docx_path, document_context):
    """
    Populates a .docx template with data from the document_context.
    Placeholders should be in the format {{CLIENT_NAME}}, {{SELLER_NAME}}, etc.
    """
    if not os.path.exists(docx_path):
        logging.error(f"DOCX template file not found: {docx_path}")
        QMessageBox.critical(None, "Template Error", f"DOCX template file not found: {docx_path}")
        return

    try:
        document = Document(docx_path)

        # Extract data from document_context
        client = document_context.get("client", {})
        seller = document_context.get("seller", {})
        doc_info = document_context.get("doc", {})
        project_info = document_context.get("project", {})
        # products_info = document_context.get("products", []) # Products might be complex for simple placeholders

        placeholders = {
            "{{CLIENT_NAME}}": client.get("individual_name", ""),
            "{{CLIENT_COMPANY_NAME}}": client.get("company_name", ""),
            "{{CLIENT_DISPLAY_NAME}}": client.get("display_name", ""),
            "{{CLIENT_ADDRESS_FULL}}": client.get("address_full", ""),
            "{{CLIENT_ADDRESS_LINE1}}": client.get("address_line1", ""),
            "{{CLIENT_CITY}}": client.get("city", ""),
            "{{CLIENT_COUNTRY}}": client.get("country", ""),
            "{{CLIENT_POSTAL_CODE}}": client.get("postal_code", ""),
            "{{CLIENT_CONTACT_NAME}}": client.get("contact_name", ""),
            "{{CLIENT_CONTACT_EMAIL}}": client.get("contact_email", ""),
            "{{CLIENT_CONTACT_PHONE}}": client.get("contact_phone", ""),
            "{{CLIENT_VAT_ID}}": client.get("vat_id", ""),
            "{{CLIENT_REGISTRATION_NUMBER}}": client.get("registration_number", ""),

            "{{SELLER_NAME}}": seller.get("name", ""),
            "{{SELLER_ADDRESS_FULL}}": seller.get("address_full", ""),
            "{{SELLER_ADDRESS_LINE1}}": seller.get("address_line1", ""),
            "{{SELLER_CITY}}": seller.get("city", ""),
            "{{SELLER_COUNTRY}}": seller.get("country", ""),
            "{{SELLER_POSTAL_CODE}}": seller.get("postal_code", ""),
            "{{SELLER_PHONE}}": seller.get("phone", ""),
            "{{SELLER_EMAIL}}": seller.get("email", ""),
            "{{SELLER_WEBSITE}}": seller.get("website", ""),
            "{{SELLER_VAT_ID}}": seller.get("vat_id", ""),
            "{{SELLER_REGISTRATION_NUMBER}}": seller.get("registration_number", ""),
            "{{SELLER_BANK_NAME}}": seller.get("bank_name", ""),
            "{{SELLER_BANK_ACCOUNT_HOLDER_NAME}}": seller.get("bank_account_holder_name", ""),
            "{{SELLER_BANK_IBAN}}": seller.get("bank_iban", ""),
            "{{SELLER_BANK_SWIFT_BIC}}": seller.get("bank_swift_bic", ""),
            "{{SELLER_BANK_ADDRESS}}": seller.get("bank_address", ""),
            "{{SELLER_REPRESENTATIVE_NAME}}": seller.get("representative_name", ""),
            "{{SELLER_LOGO_URL}}": seller.get("logo_url", ""), # Note: DOCX might not directly embed web URLs for images easily. Path might be better.

            "{{PROJECT_NAME}}": project_info.get("name", ""),
            "{{PROJECT_IDENTIFIER}}": project_info.get("identifier", ""),

            "{{DOC_CURRENT_DATE}}": doc_info.get("current_date_formatted", doc_info.get("current_date", "")),
            "{{DOC_REFERENCE}}": doc_info.get("document_reference", ""),
            "{{DOC_PROFORMA_ID}}": doc_info.get("proforma_id", ""), # Alias for proforma
            "{{DOC_CURRENCY_SYMBOL}}": doc_info.get("currency_symbol", ""),
            "{{DOC_SUBTOTAL_AMOUNT}}": doc_info.get("subtotal_amount", ""),
            "{{DOC_DISCOUNT_AMOUNT}}": doc_info.get("discount_amount", ""),
            "{{DOC_VAT_AMOUNT}}": doc_info.get("vat_amount", ""),
            "{{DOC_GRAND_TOTAL_AMOUNT}}": doc_info.get("grand_total_amount", ""),
            # Add more from doc_info.get("client_specific_notes") if it's a simple string, or handle list
        }

        # Add labels from context if they exist
        labels = document_context.get("labels", {})
        for label_key, label_value in labels.items():
            placeholders[f"{{{{LABEL_{label_key.upper()}}}}}"] = label_value # e.g. {{LABEL_PROFORMA_NO}}

        # Simple replacement for paragraphs and tables
        # For more complex scenarios (like repeating product rows), python-docx-template might be better.
        for para in document.paragraphs:
            for key, value in placeholders.items():
                if key in para.text:
                    # Replacing in runs to preserve formatting as much as possible
                    for run in para.runs:
                        if key in run.text:
                            run.text = run.text.replace(key, str(value))

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for key, value in placeholders.items():
                            if key in para.text:
                                for run in para.runs:
                                    if key in run.text:
                                        run.text = run.text.replace(key, str(value))

        document.save(docx_path)
        logging.info(f"DOCX template populated using document_context: {docx_path}")
    except Exception as e:
        logging.error(f"Error populating DOCX template {docx_path} with document_context: {e}", exc_info=True)
        QMessageBox.critical(None, "Population Error", f"Error populating DOCX template {docx_path}: {e}")
        return
